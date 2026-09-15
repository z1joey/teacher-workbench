"""成绩 Excel 批量导入：模板下载、按学号 upsert 成绩、逐行成功/失败报告，
以及导入后学生时间线出现「参加考试」（exam_taken）记录。"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, time

from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app import eventing
from app.models import AuthSession, Class, Enrollment, Event, Person
from app.payloads import validate_person_payload
from app.routers import exams as exams_router
from app.routers import students as students_router
from app.security import hash_password
from tests.test_scores_events import (
    _enroll,
    _headers,
    _seed_class,
    _seed_person,
    _seed_teacher,
)

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def ctx(make_client, db):
    teacher = _seed_teacher(db)
    headers = _headers(db, teacher)
    client = make_client(exams_router.router, students_router.router)

    cls = _seed_class(db)
    lin = _seed_person(db, "林晓雨", "S001")
    hao = _seed_person(db, "王浩", "S002")
    gone = _seed_person(db, "张离校", "S003")
    gone.payload = validate_person_payload("student", {
        **gone.payload,
        "graduated_at": "2026-01-01",
    })
    for s in (lin, hao):
        _enroll(db, s, cls)
    db.commit()

    r = client.post(
        "/api/exams",
        json={
            "name": "九月月考",
            "exam_date": "2026-09-08",
            "subjects": [
                {"subject": "math", "full_score": 120},
                {"subject": "english", "full_score": 120},
            ],
        },
        headers=headers,
    )
    assert r.status_code == 201, r.text
    exam = r.json()
    return {
        "client": client,
        "db": db,
        "headers": headers,
        "teacher": teacher,
        "exam": exam,
        "students": {"lin": lin, "hao": hao, "gone": gone},
    }


def _xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(ctx, content: bytes, filename="scores.xlsx"):
    return ctx["client"].post(
        f"/api/exams/{ctx['exam']['id']}/scores/import",
        files={"file": (filename, content, _XLSX_MIME)},
        headers=ctx["headers"],
    )


def _scores_of(db, person: Person) -> list[Event]:
    return (
        db.query(Event)
        .filter(Event.type == "score", Event.attendees.any(Person.id == person.id))
        .all()
    )


def test_template_prefills_subjects_and_students(ctx):
    r = ctx["client"].get(
        f"/api/exams/{ctx['exam']['id']}/scores/import-template",
        headers=ctx["headers"],
    )
    assert r.status_code == 200
    assert r.headers["content-type"] == _XLSX_MIME
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][:2] == ("学号", "姓名")
    assert sorted(rows[1][2:]) == ["数学(满分120)", "英语(满分120)"]
    nos = {row[0] for row in rows[2:]}
    assert {"S001", "S002"} <= nos  # 在册学生已预填学号


def test_import_creates_scores_and_exam_taken_timeline(ctx):
    lin, hao = ctx["students"]["lin"], ctx["students"]["hao"]
    content = _xlsx([
        ["学号", "姓名", "数学(满分120)", "英语(满分120)"],
        ["S001", "林晓雨", 96.5, "缺考"],
        ["S002", "王浩", 88, ""],
    ])
    r = _upload(ctx, content)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["imported_students"] == 2
    assert body["entered_scores"] == 2
    assert body["absent_scores"] == 1
    assert body["errors"] == []

    db = ctx["db"]
    db.commit()  # release the read transaction before re-querying API writes
    db.expire_all()
    lin_scores = {e.payload["subject"]: e.payload for e in _scores_of(db, lin)}
    assert lin_scores["math"]["score"] == 96.5
    assert lin_scores["math"]["max_score"] == 120
    assert lin_scores["math"]["absent"] is False
    assert lin_scores["english"]["absent"] is True and "score" not in lin_scores["english"]
    # score Event 沿用 sitting 约定：标题「考试名·科目」、日期为考试当天
    math_row = db.query(Event).filter(
        Event.type == "score", Event.title == "九月月考·math"
    ).first()
    assert math_row.start_time.date() == date(2026, 9, 8)

    # 时间线：exam_taken 汇总一行，缺考科目不出现在 scores 里；exam 场次不重复出现
    tl = ctx["client"].get(f"/api/students/{lin.id}/timeline", headers=ctx["headers"]).json()
    assert all(t["event_type"] != "exam" for t in tl)
    taken = [t for t in tl if t["event_type"] == "exam_taken"]
    assert len(taken) == 1
    assert taken[0]["payload"]["exam"] == "九月月考"
    assert taken[0]["payload"]["scores"] == {"math": 96.5}
    assert taken[0]["payload"]["absent_subjects"] == ["english"]

    # 平均分立即生效（缺考不计入）
    avg = ctx["client"].get(
        f"/api/exams/{ctx['exam']['id']}/averages", headers=ctx["headers"]
    ).json()
    math = next(s for s in avg["school"] if s["subject"] == "math")
    assert math["count"] == 2 and math["avg"] == pytest.approx((96.5 + 88) / 2, abs=0.1)
    eng = next((s for s in avg["school"] if s["subject"] == "english"), None)
    assert eng is None  # 英语只有缺考 → 没有可统计的 Entered 行


def test_reimport_overwrites_without_duplicates(ctx):
    lin = ctx["students"]["lin"]
    first = _xlsx([
        ["学号", "姓名", "数学(满分120)"],
        ["S001", "林晓雨", 90],
    ])
    assert _upload(ctx, first).json()["imported_students"] == 1
    second = _xlsx([
        ["学号", "姓名", "数学(满分120)"],
        ["S001", "林晓雨", 95],
    ])
    body = _upload(ctx, second).json()
    assert body["updated_cells"] == 1

    db = ctx["db"]
    db.commit()  # release the read transaction before re-querying API writes
    db.expire_all()
    rows = [e for e in _scores_of(db, lin) if e.payload["subject"] == "math"]
    assert len(rows) == 1 and rows[0].payload["score"] == 95
    tl = ctx["client"].get(f"/api/students/{lin.id}/timeline", headers=ctx["headers"]).json()
    assert len([t for t in tl if t["event_type"] == "exam_taken"]) == 1
    assert tl and next(t for t in tl if t["event_type"] == "exam_taken")["payload"]["scores"] == {"math": 95}


def test_import_reports_per_row_errors(ctx):
    content = _xlsx([
        ["学号", "姓名", "数学(满分120)", "英语(满分120)"],
        ["S999", "不存在", 50, ""],                # 学号不存在
        ["S001", "林晓雨", 999, "abc"],             # 全部无效 → error
        ["", "", 50, ""],                           # 缺学号
        ["S002", "王浩", 88, "缺考"],
        ["S002", "王浩", 70, ""],                   # 文件内学号重复
    ])
    body = _upload(ctx, content).json()
    by_no = {r["admission_no"]: r for r in body["rows"]}
    assert by_no["S999"]["status"] == "error"
    assert by_no["S001"]["status"] == "error"
    assert body["rows"][2]["message"] == "缺少学号"
    assert by_no["S002"]["status"] == "error"
    assert body["errors"] == [r for r in body["rows"] if r["status"] == "error"]

    # 部分成功：一格超范围、一格有效
    partial = _xlsx([
        ["学号", "姓名", "数学(满分120)", "英语(满分120)"],
        ["S001", "林晓雨", 999, 66],
    ])
    body = _upload(ctx, partial).json()
    row = body["rows"][0]
    assert row["status"] == "partial"
    assert "数学" in row["message"]
    assert row["subjects"] == ["english"]
    assert body["imported_students"] == 1


def test_import_requires_header_and_existing_exam(ctx):
    bad = _upload(ctx, _xlsx([["姓名", "成绩"], ["林晓雨", 90]]))
    assert bad.status_code == 400
    assert "表头" in bad.json()["detail"]

    missing = ctx["client"].post(
        f"/api/exams/{uuid.uuid4()}/scores/import",
        files={"file": ("s.xlsx", _xlsx([["学号", "数学"], ["S001", 1]]), _XLSX_MIME)},
        headers=ctx["headers"],
    )
    assert missing.status_code == 404


def _add(ctx, student, scores, exam_id=None):
    return ctx["client"].post(
        f"/api/exams/{exam_id or ctx['exam']['id']}/scores",
        json={"student_id": str(student.id), "scores": scores},
        headers=ctx["headers"],
    )


def test_manual_scores_create_then_upsert(ctx, db):
    """学生页「添加成绩」：录分 + 缺考 → 与导入共用 upsert，不产生重复记录。"""
    lin = ctx["students"]["lin"]
    r = _add(ctx, lin, [
        {"subject": "math", "score": 90},
        {"subject": "english", "absent": True},
    ])
    assert r.status_code == 201, r.text
    assert r.json()["scores"] == {"math": 90}
    assert r.json()["absent_subjects"] == ["english"]
    assert len(_scores_of(db, lin)) == 2
    absent_event = next(e for e in _scores_of(db, lin) if e.payload.get("absent"))
    assert absent_event.title == "九月月考·english"

    # 再录一次同科目：覆盖而非新建；缺考改成有分
    r = _add(ctx, lin, [
        {"subject": "math", "score": 95},
        {"subject": "english", "score": 100},
    ])
    assert r.status_code == 201, r.text
    db.expire_all()  # 测试会话的 identity map 里是旧 payload，强制重读
    assert len(_scores_of(db, lin)) == 2
    by_subject = {e.payload["subject"]: e.payload for e in _scores_of(db, lin)}
    assert by_subject["math"]["score"] == 95
    assert by_subject["english"]["score"] == 100
    # 时间线上的 exam_taken 也只保留一条且同步了最新分数
    taken = (
        db.query(Event)
        .filter(Event.type == "exam_taken", Event.attendees.any(Person.id == lin.id))
        .all()
    )
    assert len(taken) == 1
    assert taken[0].payload["scores"] == {"math": 95, "english": 100}
    assert taken[0].payload["absent_subjects"] == []


def test_manual_scores_validation(ctx, db):
    lin = ctx["students"]["lin"]
    hao = ctx["students"]["hao"]

    unknown = _add(ctx, lin, [{"subject": "体育", "score": 80}])
    assert unknown.status_code == 400
    assert "不是这次考试的科目" in unknown.json()["detail"]

    over = _add(ctx, lin, [{"subject": "math", "score": 121}])
    assert over.status_code == 400
    assert "0–120" in over.json()["detail"]

    missing = _add(ctx, lin, [{"subject": "math"}])
    assert missing.status_code == 400
    assert "缺考" in missing.json()["detail"]

    empty = _add(ctx, lin, [])
    assert empty.status_code == 422  # scores min_length=1

    nonexistent = _add(ctx, hao, [{"subject": "math", "score": 1}], exam_id=str(uuid.uuid4()))
    assert nonexistent.status_code == 404

    # 校验失败的请求不留任何成绩
    assert _scores_of(db, lin) == []
    assert _scores_of(db, hao) == []
