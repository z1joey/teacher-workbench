"""Student roster import/export through .xlsx only."""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from app.eventing import create_event
from app.models import AuthSession, Class, Enrollment, Event, Person
from app.payloads import validate_person_payload
from app.routers import data
from app.routers.students import _guardian_link
from app.security import hash_password
from app.unassigned import ensure_unassigned_class, is_unassigned_class
from app.workspace import ensure_workspace_id, tag_student_workspace
from tests.conftest import seed_person, seed_token

TEACHER_TOKEN = "d" * 64
AUTH = {"Authorization": f"Bearer {TEACHER_TOKEN}"}
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# 与后端导出一致的完整表头（含扩展列）
_ROSTER_HEADERS_WITH_GUARDIAN = [
    "学号", "姓名", "性别", "出生日期", "家庭住址", "监护人姓名", "监护人电话", "监护人关系",
]


def _student(db, name: str, admission_no: str, gender: str | None = None,
             birth_date: str | None = None) -> Person:
    p = Person(
        name=name,
        password_hash=hash_password("123456"),
        payload=validate_person_payload(
            "student",
            {"admission_no": admission_no, "gender": gender, "birth_date": birth_date},
        ),
    )
    db.add(p)
    db.flush()
    return p


@pytest.fixture()
def client(make_client, db):
    tc = make_client(data.router)
    teacher = seed_person(db, "chen@test.example", phone="13800000001", name="陈老师")
    seed_token(db, teacher, TEACHER_TOKEN)
    db.commit()
    return tc


def _sample_workbook(title: str = "2025级707班花名册", remarks: bool = False) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([title])
    ws.append(["学号", "姓名", "性别"] + (["备注"] if remarks else []))
    ws.append(["2025070701", "吴梓涵", "女"] + ([""] if remarks else []))
    ws.append(["2025070702", "邢宇辰", "男"] + (["流感"] if remarks else []))
    ws.append([2025070703, "卢梓扬", "男"])
    return wb


def _xlsx_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, content: bytes, **fields):
    return client.post(
        "/api/data/import/roster",
        files={"file": ("roster.xlsx", content, XLSX_MIME)},
        data={k: str(v) for k, v in fields.items()},
        headers=AUTH,
    )


def _students(db):
    return db.query(Person).filter(
        Person.payload["role"].as_string() == "student"
    ).all()


def _unassigned(db):
    return ensure_unassigned_class(db)


def test_roster_import_to_unassigned_by_default(client, db):
    res = _upload(client, _xlsx_bytes(_sample_workbook()))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 3
    assert body["updated"] == 0
    assert body["errors"] == []
    assert body["target_class"] is None

    unassigned = _unassigned(db)
    students = _students(db)
    assert {s.name for s in students} == {"吴梓涵", "邢宇辰", "卢梓扬"}
    wu = next(s for s in students if s.name == "吴梓涵")
    assert wu.payload["gender"] == "F"
    assert all(
        db.query(Enrollment).filter(
            Enrollment.person_id == s.id,
            Enrollment.class_id == unassigned.id,
            Enrollment.valid_to.is_(None),
        ).count() == 1
        for s in students
    )
    assert db.query(Event).filter(Event.type == "enrolled").count() == 3


def test_roster_import_into_existing_class(client, db):
    klass = Class(name="707班", academic_year="2025")
    db.add(klass)
    db.commit()

    res = _upload(client, _xlsx_bytes(_sample_workbook()), class_id=str(klass.id))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 3
    assert body["target_class"]["name"] == "707班"
    assert db.query(Class).filter(Class.name == "707班").count() == 1


def test_roster_import_updates_existing_without_moving_class(client, db):
    old = Class(name="七年级1班", academic_year="2025/2026")
    db.add(old)
    db.flush()
    existing = _student(db, "吴梓涵", "2025070701", gender="M")
    db.add(Enrollment(person_id=existing.id, class_id=old.id,
                      valid_from=date(2025, 9, 1)))
    db.commit()

    res = _upload(client, _xlsx_bytes(_sample_workbook()))
    assert res.status_code == 200
    body = res.json()
    assert body["created"] == 2
    assert body["updated"] == 1
    updated_row = next(r for r in body["rows"] if r["status"] == "updated")
    assert "性别" in updated_row["changes"]
    enrollment = db.query(Enrollment).filter_by(person_id=existing.id).one()
    assert enrollment.class_id == old.id and enrollment.valid_to is None
    db.refresh(existing)
    assert existing.payload["gender"] == "F"


def test_roster_import_with_class_id_moves_existing(client, db):
    old = Class(name="七年级1班", academic_year="2025/2026")
    target = Class(name="707班", academic_year="2025")
    db.add_all([old, target])
    db.flush()
    existing = _student(db, "吴梓涵", "2025070701")
    db.add(Enrollment(person_id=existing.id, class_id=old.id,
                      valid_from=date(2025, 9, 1)))
    db.commit()

    res = _upload(client, _xlsx_bytes(_sample_workbook()), class_id=str(target.id))
    assert res.status_code == 200
    enrollment = db.query(Enrollment).filter(
        Enrollment.person_id == existing.id, Enrollment.valid_to.is_(None)
    ).one()
    assert enrollment.class_id == target.id


def test_roster_reimport_updates_without_duplicates(client, db):
    first = _upload(client, _xlsx_bytes(_sample_workbook()))
    assert first.status_code == 200
    assert first.json()["created"] == 3

    second = _upload(client, _xlsx_bytes(_sample_workbook()))
    assert second.status_code == 200
    body = second.json()
    assert body["created"] == 0
    assert body["updated"] == 3
    # 资料没有任何变化时，更新行不携带 changes（前端据此省略明细表）
    assert all("changes" not in r for r in body["rows"])
    assert db.query(Person).count() == 4  # teacher + 3 students


def test_roster_import_row_errors_and_ignored_columns(client, db):
    wb = Workbook()
    ws = wb.active
    ws.append(["2025级707班花名册"])
    ws.append(["学号", "姓名", "性别", "备注"])
    ws.append(["2025070701", "吴梓涵", "女", "流感"])
    ws.append(["", "无学号", "女", ""])
    ws.append(["2025070701", "学号重复", "男", ""])
    res = _upload(client, _xlsx_bytes(wb))
    assert res.status_code == 200
    body = res.json()
    assert body["created"] == 1
    assert body["ignored_columns"] == ["备注"]
    assert {e["message"] for e in body["errors"]} == {"缺少学号", "文件内学号重复"}


def test_roster_import_rejects_bad_gender(client, db):
    wb = Workbook()
    ws = wb.active
    ws.append(["花名册"])
    ws.append(["学号", "姓名", "性别"])
    ws.append(["2025070701", "吴梓涵", "未知"])
    res = _upload(client, _xlsx_bytes(wb))
    assert res.status_code == 200
    body = res.json()
    assert body["created"] == 0
    assert body["errors"][0]["message"] == "无法识别的性别: 未知"


def test_roster_import_without_header_fails(client, db):
    wb = Workbook()
    wb.active.append(["只有一行字"])
    res = _upload(client, _xlsx_bytes(wb))
    assert res.status_code == 400
    assert "未找到表头行" in res.json()["detail"]


def test_roster_export_roundtrip(client, db):
    klass = Class(name="707班", academic_year="2025")
    db.add(klass)
    db.flush()
    s1 = _student(db, "吴梓涵", "2025070701", gender="F", birth_date="2012-05-14")
    _student(db, "邢宇辰", "2025070702", gender="M")
    payload = dict(s1.payload or {})
    payload["address"] = "解放路100号"
    s1.payload = validate_person_payload("student", payload)
    db.flush()
    guardian = Person(
        name="吴母", phone="13900000001",
        password_hash=hash_password("123456"),
        payload=validate_person_payload("guardian", {"phone": "13900000001"}),
    )
    db.add(guardian)
    db.flush()
    _guardian_link(db, s1.id, guardian, "母亲")
    for s in _students(db):
        db.add(Enrollment(person_id=s.id, class_id=klass.id, valid_from=date(2025, 9, 1)))
    db.commit()

    res = client.get(
        "/api/data/export/roster", params={"class_id": str(klass.id)}, headers=AUTH
    )
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws.cell(3, 1).value == "2025070701" and ws.cell(3, 2).value == "吴梓涵"
    assert ws.cell(3, 3).value == "女"
    assert ws.cell(4, 3).value == "男"
    # 扩展列：出生日期 / 家庭住址 / 首位监护人（姓名/电话/关系）
    assert ws.cell(3, 4).value == "2012-05-14"
    assert ws.cell(3, 5).value == "解放路100号"
    assert ws.cell(3, 6).value == "吴母"
    assert ws.cell(3, 7).value == "13900000001"
    assert ws.cell(3, 8).value == "母亲"

    res2 = _upload(client, res.content)
    assert res2.status_code == 200
    body = res2.json()
    assert body["created"] == 0
    assert body["updated"] == 2
    unassigned = _unassigned(db)
    still_in_class = db.query(Enrollment).filter(
        Enrollment.valid_to.is_(None),
        Enrollment.class_id == klass.id,
    ).count()
    assert still_in_class == 2
    assert not is_unassigned_class(klass)


def test_roster_export_unassigned_class_404(client, db):
    unassigned = _unassigned(db)
    db.commit()
    res = client.get(f"/api/data/export/roster?class_id={unassigned.id}", headers=AUTH)
    assert res.status_code == 404


def test_roster_import_links_guardian_and_merges_shared_guardian(client, db):
    """监护人列：挂接 + 关系写入；两行填同名同电话的监护人时合并为同一位。"""
    wb = Workbook()
    ws = wb.active
    ws.append(_ROSTER_HEADERS_WITH_GUARDIAN)
    ws.append(["2025070801", "张小凡", "男", "2013-03-04", "朝阳路8号", "张大力", "13900000003", "父亲"])
    ws.append(["2025070802", "张小乐", "女", "2014-07-11", "朝阳路8号", "张大力", "13900000003", "父亲"])
    res = _upload(client, _xlsx_bytes(wb))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 2
    assert all(r.get("guardian") == "张大力" for r in body["rows"])

    students = {s.name: s for s in _students(db)}
    xiaofan = students["张小凡"]
    xiaole = students["张小乐"]
    payload = xiaofan.payload or {}
    assert payload["birth_date"] == "2013-03-04"
    assert payload["address"] == "朝阳路8号"
    # 同名同电话 → 合并为同一位监护人，两个孩子都挂在他名下
    assert len(xiaofan.guardians) == 1
    assert len(xiaole.guardians) == 1
    assert xiaofan.guardians[0].id == xiaole.guardians[0].id
    assert xiaofan.guardians[0].name == "张大力"
    assert xiaofan.guardians[0].phone == "13900000003"

    # 关系写在学生-监护人关联上；再导入一次不产生重复挂接
    before = len(xiaofan.guardians)
    res2 = _upload(client, _xlsx_bytes(wb))
    assert res2.status_code == 200
    assert res2.json()["updated"] == 2
    db.expire_all()
    assert len(xiaofan.guardians) == before


def test_roster_template_download(client):
    res = client.get("/api/data/export/roster-template", headers=AUTH)
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    assert [ws.cell(2, c).value for c in range(1, 9)] == [
        "学号", "姓名", "性别", "出生日期", "家庭住址", "监护人姓名", "监护人电话", "监护人关系",
    ]
    # 模板示例行带完整字段
    assert ws.cell(3, 4).value == "2012-05-14"
    assert ws.cell(3, 6).value == "张丽"


def test_demo_seed_loads_dataset(client, db):
    res = client.post("/api/data/demo/seed", headers=AUTH)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["ok"] is True
    assert body["teacher"]["phone"] == "13800000001"
    students = _students(db)
    assert len(students) >= 20
    assert db.query(Class).filter(Class.name == "七年级1班").count() == 1
    assert db.query(Event).filter(Event.type == "exam").count() >= 1
    assert db.query(AuthSession).filter(AuthSession.token == TEACHER_TOKEN).count() == 1
    assert db.query(Person).filter(Person.phone == "13800000000").count() == 0


def test_demo_seed_requires_teacher(client, db):
    from app.security import hash_password

    student = Person(
        name="林小明",
        password_hash=hash_password("123456"),
        payload=validate_person_payload("student", {"admission_no": "S999"}),
    )
    db.add(student)
    db.flush()
    token = "s" * 64
    db.add(AuthSession(token=token, person_id=student.id))
    db.commit()

    res = client.post("/api/data/demo/seed", headers={"Authorization": f"Bearer {token}"})
    assert res.status_code == 403


def test_demo_reset_clears_database(client, db):
    seed_res = client.post("/api/data/demo/seed", headers=AUTH)
    assert seed_res.status_code == 200
    assert len(_students(db)) >= 20

    reset_res = client.post("/api/data/demo/reset", headers=AUTH)
    assert reset_res.status_code == 200
    body = reset_res.json()
    assert body["ok"] is True
    assert body["teacher"]["phone"] == "13800000001"
    assert len(_students(db)) == 0
    assert db.query(Event).count() == 0
    assert db.query(Class).filter(Class.name == "七年级1班").count() == 0


def test_roster_import_rejects_score_sheet(client, db):
    """成绩模板混进花名册导入会静默覆盖真实学生资料，必须在入口拒收。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩导入"
    ws.append(["九月月考成绩导入（2026-09-08）"])
    ws.append(["学号", "姓名", "英语(满分120)", "数学(满分120)"])
    ws.append(["2025070701", "张三", 90, 88])
    ws.append(["2025070702", "李四", 95, 91])
    res = _upload(client, _xlsx_bytes(wb))
    assert res.status_code == 400
    assert "成绩表格" in res.json()["detail"]
    assert "满分120" in res.json()["detail"]
    # 拒收不能有任何副作用：不建学生、不动账号
    assert db.query(Person).filter(
        Person.payload["role"].as_string() == "student"
    ).count() == 0
    assert db.query(AuthSession).filter(AuthSession.token == TEACHER_TOKEN).count() == 1


# ---------------------------------------------------------------------------
# 演示数据加载保护：已有业务数据时必须先「清空业务数据」
# ---------------------------------------------------------------------------

def _teacher_row(db) -> Person:
    return db.query(Person).filter(Person.phone == "13800000001").one()


def _workspace_student(db, name: str, admission_no: str) -> Person:
    teacher = _teacher_row(db)
    p = _student(db, name, admission_no)
    tag_student_workspace(p, teacher)
    db.commit()
    return p


def test_demo_status_empty_workspace(client, db):
    res = client.get("/api/data/demo/status", headers=AUTH)
    assert res.status_code == 200
    assert res.json() == {"has_business_data": False}


def test_demo_status_detects_students(client, db):
    _workspace_student(db, "林晓雨", "S770001")
    res = client.get("/api/data/demo/status", headers=AUTH)
    assert res.status_code == 200
    assert res.json() == {"has_business_data": True}


def test_demo_seed_blocked_when_students_exist(client, db):
    s = _workspace_student(db, "林晓雨", "S770001")

    res = client.post("/api/data/demo/seed", headers=AUTH)
    assert res.status_code == 409, res.text
    assert "清空业务数据" in res.json()["detail"]
    # 拒收不能有任何副作用：学生还在，演示数据没有进来
    db.expire_all()
    assert db.get(Person, s.id) is not None
    assert db.query(Class).filter(Class.name == "七年级1班").count() == 0


def test_demo_seed_blocked_when_only_exam_exists(client, db):
    """没有学生/班级、只有工作区考试时同样拦截。"""
    teacher = _teacher_row(db)
    create_event(
        db, event_type="exam", title="期末考试",
        start_time=datetime(2026, 9, 17, 9, 0),
        payload={"full_scores": {"math": 100.0},
                 "workspace_id": ensure_workspace_id(teacher)},
        attendee_ids=[teacher.id],
    )
    db.commit()

    res = client.post("/api/data/demo/seed", headers=AUTH)
    assert res.status_code == 409, res.text


def test_demo_seed_allowed_after_reset(client, db):
    _workspace_student(db, "林晓雨", "S770001")
    assert client.post("/api/data/demo/seed", headers=AUTH).status_code == 409

    assert client.post("/api/data/demo/reset", headers=AUTH).status_code == 200
    res = client.post("/api/data/demo/seed", headers=AUTH)
    assert res.status_code == 200, res.text
    assert len(_students(db)) >= 20
