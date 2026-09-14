"""Exams: a sitting is one Event(type="exam") row, scores are score Events.

The sitting Event is titled with the exam name and dated the first exam day
(start_time 09:00); multi-day sittings (中考/高考 style) put the last day in
end_time. An optional class_id scopes its attendees to that class's
enrolled students (otherwise the whole active student body attends). There
are no ExamSubject rows: the per-subject full_score config posted to
POST/PATCH /exams lives in the sitting Event's registry-validated payload as
payload["full_scores"] ({subject: full score}) and optional
payload["subject_colors"] ({subject: "#rrggbb"}) — that is where the
score-entry flow reads each subject's max_score from before writing
per-student score Events. description stays plain free text.

Score rows follow the students-router convention: one Event(type="score") per
student per subject, title "<exam name>·<subject>", start_time on the exam's
first day, payload {subject, max_score, score|absent}. With no parent link, a
sitting's scores are matched by title prefix ("{exam.title}·") AND the
sitting's date window (see sitting_score_conds) — the same rule students.py
uses to resolve exam_id. Averages aggregate payload["score"] over entered
rows only (absent=true excluded, score present).

These helpers (sitting_score_conds / subject_averages / find_exam_event) are
the shared sitting/averages vocabulary — classes.py imports them from here
rather than duplicating the aggregation.
"""
from __future__ import annotations

import io
import re
import uuid
from datetime import date, datetime, time, timedelta
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pydantic import BaseModel, Field
from sqlalchemy import and_, distinct, func, or_
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_person
from ..eventing import create_event
from ..models import Class, Enrollment, Event, Person, person_events
from ..payloads import validate_event_payload
from ..workspace import workspace_id

router = APIRouter(
    tags=["exams"],
    dependencies=[Depends(get_current_person)],  # router-level auth (include pattern)
)

EXAM_HOUR = time(9, 0)  # sitting Events are dated the exam day at 09:00


def day_window(start_day: date, end_day: date | None = None) -> tuple[datetime, datetime]:
    """Inclusive window: [00:00 on start_day, 23:59:59.999999 on end_day]
    (single-day when end_day is omitted)."""
    lo = datetime.combine(start_day, time.min)
    hi = datetime.combine(end_day or start_day, time.max)
    return lo, hi


def exam_days(e: Event) -> tuple[date, date]:
    """(first_day, last_day) of a sitting; single-day when end_time is unset."""
    return e.start_time.date(), (e.end_time or e.start_time).date()


def sitting_score_conds(title_prefix: str, start_day: date,
                        end_day: date | None = None, *,
                        wid: str) -> list:
    """Score Events of one sitting: title "<prefix>·<subject>", within the
    sitting's date window, entered only — absent=true excluded (validated
    payloads always carry absent:false explicitly, so NULL never occurs) and
    score present. `wid` keeps the title+date matching inside one teacher's
    workspace (same-named sittings in other workspaces never collide)."""
    lo, hi = day_window(start_day, end_day)
    return [
        Event.type == "score",
        Event.title.startswith(f"{title_prefix}·", autoescape=True),
        Event.start_time >= lo,
        Event.start_time < hi,
        Event.payload["workspace_id"].as_string() == wid,
        Event.payload["absent"].as_boolean().is_not(True),
        Event.payload["score"].is_not(None),
    ]


def any_sitting_score_conds(title_prefix: str, start_day: date,
                            end_day: date | None = None, *,
                            wid: str) -> list:
    """Score Events of a sitting regardless of entered state (the
    structure-frozen check in PATCH /exams mirrors the old any-ExamResult rule)."""
    lo, hi = day_window(start_day, end_day)
    return [
        Event.type == "score",
        Event.title.startswith(f"{title_prefix}·", autoescape=True),
        Event.start_time >= lo,
        Event.start_time < hi,
        Event.payload["workspace_id"].as_string() == wid,
    ]


def subject_averages(db: Session, title_prefix: str, start_day: date,
                     end_day: date | None = None,
                     person_ids: list[uuid.UUID] | None = None, *,
                     wid: str) -> dict[str, dict]:
    """Per-subject aggregate over one sitting's entered score Events:
    {subject: {avg, min, max, count, full}} — full is the max payload
    max_score (the old ExamSubject.full_score now lives on every score row).
    person_ids restricts the aggregate to those students (class attribution);
    wid restricts it to one teacher workspace."""
    subject = Event.payload["subject"].as_string()
    q = (
        db.query(
            subject,
            func.avg(Event.payload["score"].as_numeric(10, 2)),
            func.min(Event.payload["score"].as_numeric(10, 2)),
            func.max(Event.payload["score"].as_numeric(10, 2)),
            # 成绩事件可挂多名参与者（学生 + 录入老师），按事件去重计数
            func.count(distinct(Event.id)),
            func.max(Event.payload["max_score"].as_numeric(10, 2)),
        )
        .select_from(Event)
        .join(person_events, person_events.c.event_id == Event.id)
    )
    if person_ids is not None:
        q = q.filter(person_events.c.person_id.in_(person_ids))
    rows = q.filter(
        *sitting_score_conds(title_prefix, start_day, end_day, wid=wid)
    ).group_by(subject).all()
    return {
        s: {"avg": avg, "min": min_, "max": max_, "count": count, "full": full}
        for s, avg, min_, max_, count, full in rows
    }


def find_exam_event(db: Session, name: str, day: date, *, wid: str) -> Event | None:
    """The sitting Event of a name whose [first_day, last_day] contains day
    (students.py resolves score rows' exam_id with the same rule), scoped to
    one teacher workspace."""
    return (
        db.query(Event)
        .filter(
            Event.type == "exam",
            Event.title == name,
            Event.payload["workspace_id"].as_string() == wid,
            Event.start_time <= datetime.combine(day, time.max),
            func.coalesce(Event.end_time, Event.start_time)
            >= datetime.combine(day, time.min),
        )
        .first()
    )


def exam_events(db: Session, wid: str) -> list[Event]:
    """All sittings of one teacher workspace, chronological (start_time, then
    creation order — the old auto-increment id tiebreak)."""
    return (
        db.query(Event)
        .filter(
            Event.type == "exam",
            Event.payload["workspace_id"].as_string() == wid,
        )
        .order_by(Event.start_time, Event.created_at)
        .all()
    )


def _owned_exam(db: Session, user: Person, exam_id: uuid.UUID) -> Event:
    """The sitting of the caller's workspace; missing or foreign → 404 (no
    existence leak across workspaces)."""
    e = db.get(Event, exam_id)
    if (
        e is None
        or e.type != "exam"
        or (e.payload or {}).get("workspace_id") != workspace_id(user)
    ):
        raise HTTPException(status_code=404, detail="exam not found")
    return e


def subjects_config(exam: Event) -> list[dict]:
    """The sitting's subjects from the registry-validated payload
    ({"full_scores": {subject: full score}, optional "subject_colors"}),
    serialized as [{id, subject, full_score, color?}] sorted by subject —
    [] on missing or malformed config. Color is omitted when unset so older
    clients and tests keep seeing the original three-field shape. Entry ids
    are deterministic per (exam, subject) so they are stable across reads.
    """
    full_scores = (exam.payload or {}).get("full_scores")
    if not isinstance(full_scores, dict):
        return []
    colors = (exam.payload or {}).get("subject_colors")
    if not isinstance(colors, dict):
        colors = {}
    out = []
    for subject, full in sorted(full_scores.items()):
        item = {
            "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"{exam.id}:{subject}")),
            "subject": subject,
            "full_score": full,
        }
        color = colors.get(subject)
        if isinstance(color, str) and color:
            item["color"] = color
        out.append(item)
    return out


def exam_config_payload(subjects: list["SubjectIn"], term: str | None = None,
                        wid: str | None = None) -> dict:
    """Build the exam Event payload from the posted subject list. Duplicate
    names are rejected; colors are stored only when at least one is set.
    `wid` stamps the owning teacher workspace (exam isolation)."""
    names = [s.subject.strip() for s in subjects]
    if len(names) != len(set(names)):
        raise HTTPException(status_code=400, detail="科目不能重复")
    payload: dict = {"full_scores": {s.subject.strip(): s.full_score for s in subjects}}
    if term:
        payload["term"] = term
    if wid:
        payload["workspace_id"] = wid
    colors = {
        s.subject.strip(): s.color.lower()
        for s in subjects
        if s.color
    }
    if colors:
        payload["subject_colors"] = colors
    return payload


def exam_out(exam: Event) -> dict:
    # 全部学生参与者都已毕业 → 考试视为结束（前端据此归入已结束分组）
    students = [p for p in exam.attendees if (p.payload or {}).get("role") == "student"]
    return {
        "id": str(exam.id),
        "name": exam.title,
        "exam_date": exam.start_time.date().isoformat(),
        "end_date": exam.end_time.date().isoformat() if exam.end_time else None,
        "students_graduated": bool(students) and all(
            (p.payload or {}).get("graduated_at") for p in students
        ),
        "subjects": subjects_config(exam),
    }


def _attendee_ids(db: Session, user: Person,
                  class_ids: list[uuid.UUID] | None) -> list[uuid.UUID]:
    """给了班级列表就取这些班的在读学生并集；空列表 = 本工作区全部在读学生参加。"""
    if class_ids:
        return [
            row[0]
            for row in db.query(Enrollment.person_id)
            .filter(Enrollment.class_id.in_(class_ids), Enrollment.valid_to.is_(None))
            .all()
        ]
    # school-wide sitting: the owning workspace's active students attend
    active = or_(
        Person.payload["is_active"].as_boolean().is_(None),
        Person.payload["is_active"].as_boolean().is_not(False),
    )
    return [
        row[0]
        for row in db.query(Person.id)
        .filter(
            Person.payload["role"].as_string() == "student",
            Person.payload["workspace_id"].as_string() == workspace_id(user),
            active,
        )
        .all()
    ]


class SubjectIn(BaseModel):
    subject: str = Field(min_length=1, max_length=50)
    full_score: float = Field(gt=0, le=1000)
    color: str | None = Field(default=None, pattern=r"^#[0-9a-fA-F]{6}$")


class ExamIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    exam_date: date
    end_date: date | None = None  # last day of a multi-day sitting
    subjects: list[SubjectIn] = Field(min_length=1)
    class_ids: list[uuid.UUID] = Field(default_factory=list)  # 按班级圈定参加者；空 = 全校
    term: str | None = None


@router.post("/exams", status_code=201)
def create_exam(
    body: ExamIn,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    name = body.name.strip()
    if body.end_date is not None and body.end_date < body.exam_date:
        raise HTTPException(status_code=400, detail="结束日期不能早于考试日期")
    wid = workspace_id(user)
    # duplicate rule: same-named exam overlapping the [start, end] span
    # (scoped to the caller's workspace — other workspaces never collide)
    lo, hi = day_window(body.exam_date, body.end_date)
    overlap = (
        db.query(Event.id)
        .filter(
            Event.type == "exam",
            Event.title == name,
            Event.payload["workspace_id"].as_string() == wid,
            Event.start_time <= hi,
            func.coalesce(Event.end_time, Event.start_time) >= lo,
        )
        .first()
    )
    if overlap is not None:
        raise HTTPException(status_code=409, detail="该日期已存在同名考试")
    # 按班级圈定参加者：去重、校验班级存在且属于本工作区；名单取各班当前在读学生并集
    class_ids = list(dict.fromkeys(body.class_ids))
    for cid in class_ids:
        cls = db.get(Class, cid)
        if cls is None or cls.teacher_id != user.id:
            raise HTTPException(status_code=400, detail="class not found")
    exam = create_event(
        db,
        event_type="exam",
        title=name,
        start_time=datetime.combine(body.exam_date, EXAM_HOUR),
        end_time=datetime.combine(body.end_date, time.max) if body.end_date else None,
        payload=exam_config_payload(body.subjects, body.term, wid),
        # the sitting involves the teacher arranging it plus its students
        attendee_ids=[user.id, *_attendee_ids(db, user, class_ids)],
    )
    db.commit()
    return {"id": str(exam.id), "name": exam.title,
            "exam_date": exam.start_time.date().isoformat(),
            "end_date": exam.end_time.date().isoformat() if exam.end_time else None}


@router.get("/exams")
def list_exams(
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    exams = exam_events(db, workspace_id(user))
    exams.reverse()  # old listing was exam_date desc
    return [exam_out(e) for e in exams]


@router.get("/exams/trend")
def exams_trend(db: Session = Depends(get_db), user: Person = Depends(get_current_person)):
    """School-wide per-subject averages across sittings (entered scores only)."""
    wid = workspace_id(user)
    exams = exam_events(db, wid)
    index_of = {e.id: i for i, e in enumerate(exams)}
    per_subject: dict[str, dict] = {}
    for e in exams:
        start_day, end_day = exam_days(e)
        for subject, agg in subject_averages(
            db, e.title, start_day, end_day, wid=wid
        ).items():
            rec = per_subject.setdefault(
                subject, {"full_score": 0.0, "values": [None] * len(exams)}
            )
            rec["full_score"] = max(rec["full_score"], float(agg["full"] or 0))
            if e.id in index_of and agg["avg"] is not None:
                rec["values"][index_of[e.id]] = round(float(agg["avg"]), 1)
    return {
        "exams": [
            {
                "id": str(e.id),
                "name": e.title,
                "exam_date": e.start_time.date().isoformat(),
            }
            for e in exams
        ],
        "series": [
            {"subject": subject, "values": rec["values"], "full_score": rec["full_score"]}
            for subject, rec in sorted(per_subject.items())
        ],
    }


@router.get("/exams/{exam_id}")
def get_exam(
    exam_id: uuid.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    e = _owned_exam(db, user, exam_id)
    # 参加班级回显：考试学生参与者当前在读班级的并集
    student_ids = [
        p.id for p in e.attendees if (p.payload or {}).get("role") == "student"
    ]
    class_ids: list = []
    if student_ids:
        class_ids = [
            row[0]
            for row in db.query(Enrollment.class_id)
            .join(Person, Person.id == Enrollment.person_id)
            .filter(
                Enrollment.person_id.in_(student_ids),
                Enrollment.valid_to.is_(None),
            )
            .distinct()
            .all()
        ]
    return {**exam_out(e), "class_ids": class_ids}


@router.get("/exams/{exam_id}/averages")
def exam_averages(
    exam_id: uuid.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    exam = _owned_exam(db, user, exam_id)
    wid = workspace_id(user)
    start_day, end_day = exam_days(exam)

    school = [
        {
            "subject": subject,
            "full_score": float(agg["full"]) if agg["full"] is not None else None,
            "avg": round(float(agg["avg"]), 1) if agg["avg"] is not None else None,
            "min": round(float(agg["min"]), 1) if agg["min"] is not None else None,
            "max": round(float(agg["max"]), 1) if agg["max"] is not None else None,
            "count": agg["count"],
        }
        for subject, agg in sorted(
            subject_averages(db, exam.title, start_day, end_day, wid=wid).items()
        )
    ]

    # per-class averages attribute each score to the class roster valid at the
    # exam date (the old enrollment-valid-at-exam_date rule; attribution uses
    # the first day — scores follow the sitting window)
    class_rows = (
        db.query(
            Class.id,
            Class.name,
            Event.payload["subject"].as_string(),
            func.avg(Event.payload["score"].as_numeric(10, 2)),
            func.count(Event.id),
        )
        .select_from(Event)
        .join(person_events, person_events.c.event_id == Event.id)
        .join(Person, Person.id == person_events.c.person_id)
        .join(
            Enrollment,
            and_(
                Enrollment.person_id == Person.id,
                Enrollment.valid_from <= start_day,
                or_(Enrollment.valid_to.is_(None), Enrollment.valid_to >= start_day),
            ),
        )
        .join(Class, Class.id == Enrollment.class_id)
        .filter(*sitting_score_conds(exam.title, start_day, end_day, wid=wid))
        .group_by(Class.id, Class.name, Event.payload["subject"].as_string())
        .order_by(Class.name, Event.payload["subject"].as_string())
        .all()
    )

    return {
        "exam": {
            "id": str(exam.id),
            "name": exam.title,
            "exam_date": start_day.isoformat(),
            "end_date": end_day.isoformat() if exam.end_time else None,
        },
        "school": school,
        "classes": [
            {
                "class_id": str(class_id),
                "class_name": class_name,
                "subject": subject,
                "avg": round(float(avg), 1) if avg is not None else None,
                "count": count,
            }
            for class_id, class_name, subject, avg, count in class_rows
        ],
    }


class ExamUpdateIn(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=100)
    exam_date: date | None = None
    end_date: date | None = None  # last day of a multi-day sitting
    subjects: list[SubjectIn] | None = None
    class_ids: list[uuid.UUID] | None = None  # absent = 不修改参加班级


@router.patch("/exams/{exam_id}")
def update_exam(
    exam_id: uuid.UUID,
    body: ExamUpdateIn,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    e = _owned_exam(db, user, exam_id)
    wid = workspace_id(user)

    old_name, (old_start, old_end) = e.title, exam_days(e)
    new_start = body.exam_date or old_start
    # end_date present-but-null clears the span (back to a single day);
    # absent means "no change" (PATCH semantics) — except a single-day
    # sitting has no explicit end, so its implied end follows the start
    if "end_date" in body.model_fields_set:
        new_end = body.end_date or new_start
    elif e.end_time is None:
        new_end = new_start
    else:
        new_end = old_end
    if new_end < new_start:
        raise HTTPException(status_code=400, detail="结束日期不能早于考试日期")

    if body.subjects is not None:
        if (
            db.query(Event.id)
            .filter(*any_sitting_score_conds(old_name, old_start, old_end, wid=wid))
            .first()
            is not None
        ):
            raise HTTPException(status_code=400, detail="考试已有成绩录入，无法修改科目结构")
        # copy-modify-reassign + re-validate: JSON columns don't see in-place
        # mutation, and the config must stay registry-shaped
        payload = dict(e.payload or {})
        payload.pop("full_scores", None)
        payload.pop("subject_colors", None)
        payload.update(exam_config_payload(
            body.subjects, payload.get("term"), payload.get("workspace_id")
        ))
        e.payload = validate_event_payload("exam", payload)

    if body.name is not None:
        e.title = body.name.strip()
    if body.exam_date is not None:
        e.start_time = datetime.combine(body.exam_date, e.start_time.time())
    if body.end_date is not None:
        e.end_time = datetime.combine(body.end_date, time.max)
    elif "end_date" in body.model_fields_set:
        e.end_time = None  # explicitly cleared: single-day sitting again

    # renaming / re-dating the sitting must not orphan its scores: the results
    # follow the exam (the old FK behavior — results stayed attached); scores
    # are dated on the first day, so only a start-day shift moves them
    if e.title != old_name or e.start_time.date() != old_start:
        shift = (e.start_time.date() - old_start).days
        for score in (
            db.query(Event)
            .filter(*any_sitting_score_conds(old_name, old_start, old_end, wid=wid))
            .all()
        ):
            subject = score.title.rsplit("·", 1)[-1]
            score.title = f"{e.title}·{subject}"
            score.start_time = score.start_time + timedelta(days=shift)

    # 参加班级：按所选班级的当前在读学生重建学生参与者（教师保留）
    if "class_ids" in body.model_fields_set:
        class_ids = list(dict.fromkeys(body.class_ids))
        for cid in class_ids:
            cls = db.get(Class, cid)
            if cls is None or cls.teacher_id != user.id:
                raise HTTPException(status_code=400, detail="class not found")
        keep = {p.id for p in e.attendees if (p.payload or {}).get("role") != "student"}
        keep.add(user.id)
        students = _attendee_ids(db, user, class_ids)
        e.attendees = db.query(Person).filter(
            Person.id.in_([user.id, *students, *keep])
        ).all()

    db.commit()
    db.refresh(e)
    return exam_out(e)


@router.delete("/exams/{exam_id}")
def delete_exam(
    exam_id: uuid.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Deletes the sitting Event only. Score events are individual rows with
    no parent link, so they deliberately survive the delete (the old cascade
    over exam_subject/exam_result has no equivalent to walk)."""
    e = _owned_exam(db, user, exam_id)
    db.delete(e)
    db.commit()
    return {"ok": True}


# --- 成绩 Excel 批量导入 -----------------------------------------------------
#
# 模板与导入共用一套列头约定：学号、姓名（仅作人工核对），之后每门科目一列，
# 列头写「数学(满分120)」或直接「数学」。单元格里数字为成绩，「缺考」记为缺考，
# 留空跳过该科目。导入按学号匹配学生，逐格 upsert 成绩 Events（与手动录入
# 同一约定），并为每位有成绩的学生维护一条 exam_taken 时间线事件（同一场
# 考试重复导入时原地更新，不产生重复行）。
#
# 考试 payload 里的科目是稳定 key（math/english/…，与前端 COMMON_SUBJECTS
# 一致），界面显示中文名。模板列头用中文名；导入时中文名与 key 都能对上，
# 但落库一律归一化回 key，平均分才不会按两个名字分组成两组。

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ABSENT_WORDS = {"缺考", "旷考", "absent"}
_SUBJECT_FULL_RE = re.compile(r"^(.*?)[\(（]\s*满分\s*([0-9.]+)\s*[\)）]\s*$")
_SUBJECT_LABELS = {
    "chinese": "语文",
    "math": "数学",
    "english": "英语",
    "politics": "道德与法治",
    "history": "历史",
    "geography": "地理",
    "biology": "生物",
    "physics": "物理",
    "chemistry": "化学",
}


def subject_label(key: str) -> str:
    return _SUBJECT_LABELS.get(key, key)


def _subject_column_alias(cell: str, subjects: list[str]) -> str | None:
    """Header cell → the sitting subject key it refers to, or None."""
    bare = _SUBJECT_FULL_RE.match(cell)
    name = (bare.group(1) if bare else cell).strip()
    if name in subjects:
        return name
    for key in subjects:
        if subject_label(key) == name:
            return key
    return None


def _cell_str(value) -> str:
    """Excel cells arrive typed: 2025070701 as float, scores as int/float."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _full_label(full) -> str:
    return str(int(full)) if isinstance(full, float) and full.is_integer() else str(full)


def _load_xlsx_rows(content: bytes) -> list[list]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析该 Excel 文件，请提供 .xlsx 格式")
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _locate_score_header(
    rows: list[list], subjects: list[str]
) -> tuple[int, int, str | None, dict[str, int]]:
    """Find the header row: 学号 column plus at least one subject column.
    Returns (row_index, admission_no_column, name_column, {subject: column})."""
    for idx, row in enumerate(rows[:10]):
        cells = [_cell_str(c) for c in row]
        no_col = cells.index("学号") if "学号" in cells else None
        if no_col is None:
            continue
        name_col = cells.index("姓名") if "姓名" in cells else None
        columns: dict[str, int] = {}
        for col, cell in enumerate(cells):
            if col == no_col or col == name_col or not cell:
                continue
            key = _subject_column_alias(cell, subjects)
            if key is not None and key not in columns:
                columns[key] = col
        if columns:
            return idx, no_col, name_col, columns
    raise HTTPException(
        status_code=400,
        detail="未找到表头行（需包含「学号」列和至少一门考试科目列，如「数学(满分120)」）",
    )


def _active_students(db: Session, wid: str) -> list[Person]:
    active = or_(
        Person.payload["is_active"].as_boolean().is_(None),
        Person.payload["is_active"].as_boolean().is_not(False),
    )
    return (
        db.query(Person)
        .filter(
            Person.payload["role"].as_string() == "student",
            Person.payload["workspace_id"].as_string() == wid,
            active,
        )
        .order_by(Person.payload["admission_no"].as_string())
        .all()
    )


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    quoted = quote(filename)
    return StreamingResponse(
        buf,
        media_type=_XLSX_MIME,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"},
    )


@router.get("/exams/{exam_id}/scores/import-template")
def score_import_template(
    exam_id: uuid.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Import template pre-filled with the sitting's subjects and the active
    students' 学号/姓名 — teachers only type the score cells."""
    exam = _owned_exam(db, user, exam_id)
    subjects = subjects_config(exam)
    if not subjects:
        raise HTTPException(status_code=400, detail="该考试没有科目配置，无法生成模板")

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩导入"
    ws.append([f"{exam.title}成绩导入（{exam.start_time.date().isoformat()}）"])
    headers = ["学号", "姓名"] + [
        f"{subject_label(s['subject'])}(满分{_full_label(s['full_score'])})" for s in subjects
    ]
    ws.append(headers)
    for s in _active_students(db, workspace_id(user)):
        ws.append([(s.payload or {}).get("admission_no") or "", s.name] + [""] * len(subjects))
    bold = Font(bold=True)
    ws.cell(row=1, column=1).font = bold
    for col in range(1, len(headers) + 1):
        ws.cell(row=2, column=col).font = bold
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 12
    return _xlsx_response(wb, f"{exam.title}成绩导入模板.xlsx")


def _upsert_score_event(
    db: Session,
    exam: Event,
    subject_name: str,
    full_score: float,
    student: Person,
    score: float | None,
    absent: bool,
    entered_by: Person | None = None,
) -> bool:
    """Upsert one per-subject score Event; returns True when an existing row
    was updated (re-import overwrites) rather than created. `entered_by` 是
    录入老师，作为事件参与者记入审计（教学足迹「录入成绩」据此统计）。"""
    lo, hi = day_window(*exam_days(exam))
    existing = (
        db.query(Event)
        .filter(
            Event.type == "score",
            Event.title == f"{exam.title}·{subject_name}",
            Event.start_time >= lo,
            Event.start_time < hi,
            Event.attendees.any(Person.id == student.id),
        )
        .first()
    )
    payload: dict = {"subject": subject_name, "max_score": full_score}
    if absent:
        payload["absent"] = True
    else:
        payload["score"] = score
        payload["absent"] = False
    # 成绩归属考试所在的工作区（隔离后跨教师同名同日不会串数据）
    payload["workspace_id"] = (exam.payload or {}).get("workspace_id")
    payload = validate_event_payload("score", payload)
    attendees = [student.id, *([entered_by.id] if entered_by else [])]
    if existing is not None:
        existing.payload = payload
        return True
    create_event(
        db,
        event_type="score",
        title=f"{exam.title}·{subject_name}",
        start_time=datetime.combine(exam.start_time.date(), EXAM_HOUR),
        payload=payload,
        attendee_ids=attendees,
    )
    return False


def _upsert_exam_taken(
    db: Session,
    exam: Event,
    student: Person,
    scores: dict[str, float],
    absent_subjects: list[str],
) -> None:
    """The student's per-sitting timeline row (type exam_taken renders as
    「参加考试 · 九月月考 — 数学 90, 英语 85」); one per (student, sitting)."""
    existing = (
        db.query(Event)
        .filter(
            Event.type == "exam_taken",
            Event.title == exam.title,
            Event.attendees.any(Person.id == student.id),
        )
        .first()
    )
    payload = validate_event_payload(
        "exam_taken",
        {"exam": exam.title, "scores": scores, "absent_subjects": absent_subjects},
    )
    if existing is not None:
        existing.payload = payload
        return
    create_event(
        db,
        event_type="exam_taken",
        title=exam.title,
        start_time=datetime.combine(exam.start_time.date(), EXAM_HOUR),
        payload=payload,
        attendee_ids=[student.id],
    )


class StudentScoreIn(BaseModel):
    subject: str = Field(min_length=1, max_length=50)
    score: float | None = Field(default=None, ge=0, le=1000)
    absent: bool = False


class StudentScoresIn(BaseModel):
    student_id: uuid.UUID
    scores: list[StudentScoreIn] = Field(min_length=1)


@router.post("/exams/{exam_id}/scores", status_code=201)
def add_student_scores(
    exam_id: uuid.UUID,
    body: StudentScoresIn,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Manually record one student's scores for a sitting (the student page's
    添加成绩 dialog) — same upsert path as the Excel import."""
    exam = _owned_exam(db, user, exam_id)
    full_by_subject = {s["subject"]: s["full_score"] for s in subjects_config(exam)}
    if not full_by_subject:
        raise HTTPException(status_code=400, detail="该考试没有科目配置，无法录入成绩")
    student = db.get(Person, body.student_id)
    if student is None or (student.payload or {}).get("role") != "student":
        raise HTTPException(status_code=404, detail="student not found")

    seen: set[str] = set()
    scores: dict[str, float] = {}
    absent_subjects: list[str] = []
    for item in body.scores:
        if item.subject not in full_by_subject:
            raise HTTPException(status_code=400, detail=f"「{item.subject}」不是这次考试的科目")
        if item.subject in seen:
            raise HTTPException(status_code=400, detail=f"科目「{item.subject}」重复")
        seen.add(item.subject)
        if item.absent:
            absent_subjects.append(item.subject)
            _upsert_score_event(
                db, exam, item.subject, full_by_subject[item.subject], student, None, True,
                entered_by=user,
            )
            continue
        if item.score is None:
            raise HTTPException(
                status_code=400,
                detail=f"「{item.subject}」没有分数：请填写分数，或勾选缺考",
            )
        if item.score > full_by_subject[item.subject]:
            raise HTTPException(
                status_code=400,
                detail=f"「{item.subject}」成绩需在 0–{_full_label(full_by_subject[item.subject])} 之间",
            )
        scores[item.subject] = item.score
        _upsert_score_event(
            db, exam, item.subject, full_by_subject[item.subject], student, item.score, False,
            entered_by=user,
        )
    if not scores and not absent_subjects:
        raise HTTPException(status_code=400, detail="没有可录入的成绩")

    _upsert_exam_taken(db, exam, student, scores, absent_subjects)
    db.commit()
    return {
        "ok": True,
        "entered": len(scores) + len(absent_subjects),
        "scores": scores,
        "absent_subjects": absent_subjects,
    }


@router.post("/exams/{exam_id}/scores/import")
async def import_scores(
    exam_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Bulk-import score rows from xlsx, matched by 学号. Every row reports
    success (with imported subjects) or failure (with reason)."""
    exam = _owned_exam(db, user, exam_id)
    subjects = subjects_config(exam)
    if not subjects:
        raise HTTPException(status_code=400, detail="该考试没有科目配置，无法导入成绩")
    full_by_subject = {s["subject"]: s["full_score"] for s in subjects}

    rows = _load_xlsx_rows(await file.read())
    header_idx, no_col, name_col, columns = _locate_score_header(rows, list(full_by_subject))
    ignored = sorted(
        {
            _cell_str(c)
            for row in rows[header_idx : header_idx + 1]
            for col, c in enumerate(row)
            if _cell_str(c) and col not in columns.values() and col not in {no_col, name_col}
        }
    )

    students_by_no = {
        (s.payload or {}).get("admission_no"): s
        for s in _active_students(db, workspace_id(user))
    }
    report: list[dict] = []
    seen_rows: set[str] = set()
    imported_students = entered = absent_cells = updated_cells = 0
    for pos, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        admission_no = _cell_str(row[no_col]) if no_col < len(row) else ""
        name = _cell_str(row[name_col]) if name_col is not None and name_col < len(row) else ""
        if not admission_no and not any(
            col < len(row) and _cell_str(row[col]) for col in columns.values()
        ):
            continue  # 整行为空
        item = {
            "row": pos,
            "admission_no": admission_no,
            "name": name,
            "status": "ok",
            "subjects": [],
            "absent_subjects": [],
            "message": None,
        }
        try:
            if not admission_no:
                raise ValueError("缺少学号")
            if admission_no in seen_rows:
                raise ValueError("文件内学号重复")
            seen_rows.add(admission_no)
            student = students_by_no.get(admission_no)
            if student is None:
                raise ValueError("学号不存在")

            warnings: list[str] = []
            scores: dict[str, float] = {}
            absent_subjects: list[str] = []
            for subject_name, col in columns.items():
                raw = _cell_str(row[col]) if col < len(row) else ""
                if not raw:
                    continue
                low = raw.lower()
                if low in _ABSENT_WORDS:
                    updated = _upsert_score_event(
                        db, exam, subject_name, full_by_subject[subject_name],
                        student, None, True, entered_by=user,
                    )
                    absent_subjects.append(subject_name)
                    absent_cells += 1
                    updated_cells += 1 if updated else 0
                    continue
                try:
                    score = float(raw)
                except ValueError:
                    warnings.append(f"{subject_label(subject_name)}: 无法识别的成绩「{raw}」")
                    continue
                if score < 0 or score > full_by_subject[subject_name]:
                    warnings.append(
                        f"{subject_label(subject_name)}: 成绩需在 0–{_full_label(full_by_subject[subject_name])} 之间"
                    )
                    continue
                updated = _upsert_score_event(
                    db, exam, subject_name, full_by_subject[subject_name],
                    student, score, False, entered_by=user,
                )
                scores[subject_name] = score
                item["subjects"].append(subject_name)
                entered += 1
                updated_cells += 1 if updated else 0
            if not scores and not absent_subjects:
                raise ValueError("；".join(warnings) or "该行没有可导入的成绩")
            if warnings:
                item["status"] = "partial"
                item["message"] = "；".join(warnings)
            _upsert_exam_taken(db, exam, student, scores, absent_subjects)
            item["absent_subjects"] = absent_subjects
            imported_students += 1
        except ValueError as exc:
            item["status"] = "error"
            item["message"] = str(exc)
        report.append(item)

    errors = [r for r in report if r["status"] == "error"]
    db.commit()
    return {
        "exam": exam_out(exam),
        "imported_students": imported_students,
        "entered_scores": entered,
        "absent_scores": absent_cells,
        "updated_cells": updated_cells,
        "errors": errors,
        "ignored_columns": ignored,
        "rows": report,
    }
