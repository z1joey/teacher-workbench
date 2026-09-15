"""Student data import/export through .xlsx.

Roster (POST /data/import/roster, GET /data/export/roster[/-template]) moves
student rows through .xlsx in the format schools actually hand out: an optional
merged title row ("…班花名册"), a header row (学号/姓名/性别[/出生日期/家庭住址/
监护人姓名/监护人电话/监护人关系]) and one student per row. Import creates or
updates students only — it never creates classes. New rows land in the system
unassigned class unless an class_id is supplied; re-import without class_id
updates fields in place without moving enrollments. Guardian columns link (or
merge, by phone/name) the guardian onto the student; they never unlink an
existing guardian.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import uuid as uuid_mod
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import delete, or_, select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_person
from ..eventing import create_event, sync_birthday_event
from ..gender import gender_label, parse_gender
from ..models import Class, ClassSeating, Enrollment, Event, Feedback, Person, Tag
from ..models._common import utcnow
from ..models.associations import person_events, person_tags, student_guardians
from ..payloads import validate_person_payload
from ..security import hash_password
from ..seed import seed
from ..unassigned import (
    UNASSIGNED_ACADEMIC_YEAR,
    UNASSIGNED_CLASS_NAME,
    class_export_heading,
    class_for_api,
    ensure_unassigned_class,
    is_unassigned_class,
)
from ..workspace import (
    classes_query,
    students_query,
    tag_student_workspace,
    workspace_id,
)
from .exams import (
    append_score_sheet,
    exam_events,
    exam_relevant_to_class,
    scores_for_students,
    subjects_config,
    unique_sheet_title,
)
from .students import (
    _find_or_create_guardian,
    _guardian_link,
    _guardians_of,
    _prune_unused_tags,
)

router = APIRouter(
    tags=["data"],
    dependencies=[Depends(get_current_person)],
)

_bearer = HTTPBearer(auto_error=False)

_HEADER_ALIASES = {
    "admission_no": {"学号", "学籍号", "学籍编号", "学籍辅号", "学籍"},
    "name": {"姓名", "名字", "学生姓名"},
    "gender": {"性别"},
    "birth_date": {"出生日期", "出生年月", "生日"},
    "address": {"家庭住址", "家庭地址", "住址", "地址"},
    "guardian_name": {"监护人姓名", "监护人", "家长姓名", "家长"},
    "guardian_phone": {"监护人电话", "监护人手机", "监护人手机号", "家长电话", "家长手机号"},
    "guardian_relation": {"监护人关系", "与学生关系", "家长关系", "关系"},
}
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# 导入/导出/模板共用同一套列，顺序即导出顺序
_ROSTER_HEADERS = [
    "学号", "姓名", "性别", "出生日期", "家庭住址", "监护人姓名", "监护人电话", "监护人关系",
]
_HOME_VISIT_HEADERS = ["学号", "姓名", "家访时间", "事由", "内容", "监护人", "完成情况"]
# 成绩表（如「英语(满分120)」）混进花名册导入会静默改掉真实学生资料，直接拒收
_SCORE_COLUMN_HINT = re.compile(r"满分|成绩|分数|得分|绩点")


def _cell_str(value) -> str:
    """Excel cells arrive typed: 2025070701 as float, dates as datetime."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_birth_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y%m%d"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法识别的出生日期: {text}")


def _locate_header(rows) -> tuple[int, dict[str, int]]:
    """Find the header row within the first 10 rows and map field -> column."""
    for idx, row in enumerate(rows[:10]):
        cells = [_cell_str(c) for c in row]
        mapping = {}
        for field, aliases in _HEADER_ALIASES.items():
            for col, cell in enumerate(cells):
                if cell in aliases:
                    mapping[field] = col
                    break
        if "admission_no" in mapping and "name" in mapping:
            return idx, mapping
    raise HTTPException(status_code=400, detail="未找到表头行（需同时包含“学号”和“姓名”列）")


def _resolve_target_class(db: Session, class_id: uuid_mod.UUID | None, user: Person) -> Class:
    if class_id is None:
        return ensure_unassigned_class(db)
    cls = db.get(Class, class_id)
    # 工作区班级只允许归属教师导入；无主班级（历史数据/夹具）保持兼容。
    if cls is None or is_unassigned_class(cls):
        raise HTTPException(status_code=404, detail="class not found")
    if cls.teacher_id is not None and cls.teacher_id != user.id:
        raise HTTPException(status_code=404, detail="class not found")
    return cls


def _class_for_export(db: Session, user: Person, class_id: uuid_mod.UUID) -> Class:
    """Export endpoints: allow 未分班; real classes need workspace ownership."""
    cls = db.get(Class, class_id)
    if cls is None:
        raise HTTPException(status_code=404, detail="class not found")
    if is_unassigned_class(cls):
        return cls
    if cls.teacher_id is not None and cls.teacher_id != user.id:
        raise HTTPException(status_code=404, detail="class not found")
    return cls


def _class_roster_students(db: Session, cls: Class) -> list[Person]:
    return (
        db.query(Person)
        .join(Enrollment, Enrollment.person_id == Person.id)
        .filter(Enrollment.class_id == cls.id, Enrollment.valid_to.is_(None))
        .order_by(Person.payload["admission_no"].as_string())
        .all()
    )


def _move_student(
    db: Session,
    person_row: Person,
    target: Class,
    *,
    reason: str,
    today: dt.date,
) -> None:
    enrollment = (
        db.query(Enrollment)
        .filter(Enrollment.person_id == person_row.id, Enrollment.valid_to.is_(None))
        .first()
    )
    current_id = enrollment.class_id if enrollment else None
    if current_id == target.id:
        return
    if enrollment is not None:
        enrollment.valid_to = today
        enrollment.reason = "moved"
    db.add(
        Enrollment(
            person_id=person_row.id,
            class_id=target.id,
            valid_from=today,
            reason=reason,
        )
    )


def _require_teacher(user: Person) -> None:
    if not user.is_teacher:
        raise HTTPException(status_code=403, detail="仅教师账号可使用演示数据功能")


def _workspace_event_ids(
    db: Session, teacher: Person, student_ids: set[uuid_mod.UUID]
) -> set[uuid_mod.UUID]:
    """Events owned by or attached to this teacher's workspace."""
    wid = workspace_id(teacher)
    exam_score_ids = {
        row[0]
        for row in db.query(Event.id).filter(
            Event.type.in_(("exam", "score")),
            Event.payload["workspace_id"].as_string() == wid,
        )
    }
    if not student_ids:
        return exam_score_ids
    linked_ids = {
        row[0]
        for row in db.execute(
            select(person_events.c.event_id).where(
                person_events.c.person_id.in_(student_ids)
            )
        )
    }
    return exam_score_ids | linked_ids


def _clear_workspace_data(
    db: Session,
    teacher: Person,
    *,
    load_seed: bool = False,
) -> Person:
    """Delete only the current teacher's workspace rows.

    Other teachers, admins, and their sessions are never touched.
    """
    student_ids = {s.id for s in students_query(db, teacher).all()}
    class_ids = {c.id for c in classes_query(db, teacher, include_archived=True).all()}
    event_ids = _workspace_event_ids(db, teacher, student_ids)

    db.execute(delete(Feedback).where(Feedback.person_id == teacher.id))

    if student_ids:
        guardian_ids = {
            row[0]
            for row in db.execute(
                select(student_guardians.c.guardian_id).where(
                    student_guardians.c.student_id.in_(student_ids)
                )
            )
        }
        db.execute(
            delete(person_tags).where(person_tags.c.person_id.in_(student_ids))
        )
        db.execute(
            delete(student_guardians).where(
                student_guardians.c.student_id.in_(student_ids)
            )
        )
    else:
        guardian_ids = set()

    if class_ids:
        db.execute(
            delete(ClassSeating).where(ClassSeating.class_id.in_(class_ids))
        )

    enrollment_filters = []
    if student_ids:
        enrollment_filters.append(Enrollment.person_id.in_(student_ids))
    if class_ids:
        enrollment_filters.append(Enrollment.class_id.in_(class_ids))
    if enrollment_filters:
        db.execute(delete(Enrollment).where(or_(*enrollment_filters)))

    if event_ids or student_ids:
        pe_filters = []
        if event_ids:
            pe_filters.append(person_events.c.event_id.in_(event_ids))
        if student_ids:
            pe_filters.append(person_events.c.person_id.in_(student_ids))
        db.execute(delete(person_events).where(or_(*pe_filters)))
    if event_ids:
        db.execute(delete(Event).where(Event.id.in_(event_ids)))

    db.execute(
        delete(Class).where(
            Class.teacher_id == teacher.id,
            ~(
                (Class.name == UNASSIGNED_CLASS_NAME)
                & (Class.academic_year == UNASSIGNED_ACADEMIC_YEAR)
            ),
        )
    )

    if student_ids:
        db.execute(delete(Person).where(Person.id.in_(student_ids)))

    if guardian_ids:
        still_linked = {
            row[0]
            for row in db.execute(
                select(student_guardians.c.guardian_id).where(
                    student_guardians.c.guardian_id.in_(guardian_ids)
                )
            )
        }
        orphan_guardian_ids = guardian_ids - still_linked
        if orphan_guardian_ids:
            role = Person.payload["role"].as_string()
            db.execute(
                delete(Person).where(
                    Person.id.in_(orphan_guardian_ids),
                    role == "guardian",
                )
            )

    _prune_unused_tags(db)
    db.flush()
    ensure_unassigned_class(db)

    kept = db.get(Person, teacher.id)
    if kept is None:
        raise RuntimeError("teacher row missing after workspace clear")
    if load_seed:
        seed(db, teacher=kept, include_admin=False)
    return kept


def _has_business_data(db: Session, teacher: Person) -> bool:
    """Whether this teacher's workspace already has roster / class / exam data."""
    if students_query(db, teacher).first() is not None:
        return True
    if classes_query(db, teacher, include_archived=True).first() is not None:
        return True
    wid = workspace_id(teacher)
    return (
        db.query(Event.id)
        .filter(
            Event.type == "exam",
            Event.payload["workspace_id"].as_string() == wid,
        )
        .first()
        is not None
    )


@router.get("/data/demo/status")
def demo_data_status(
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Whether the workspace already has business data (real or demo)."""
    _require_teacher(user)
    return {"has_business_data": _has_business_data(db, user)}


@router.post("/data/demo/seed")
def load_demo_data(
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
    credentials: HTTPAuthorizationCredentials | None = Depends(_bearer),
):
    """Wipe app data and load demo content bound to the current teacher."""
    _require_teacher(user)
    if credentials is None:
        raise HTTPException(status_code=401, detail="未登录")
    if _has_business_data(db, user):
        raise HTTPException(
            status_code=409,
            detail="当前工作台已有数据，加载演示数据前请先「清空业务数据」",
        )
    try:
        teacher = _clear_workspace_data(db, user, load_seed=True)
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"加载演示数据失败: {exc}") from exc
    return {
        "ok": True,
        "teacher": {"name": teacher.name, "phone": teacher.phone},
    }


@router.post("/data/demo/reset")
def reset_app_data(
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
    credentials: HTTPAuthorizationCredentials | None = Depends(_bearer),
):
    """Clear this teacher's workspace data but keep the account signed in."""
    _require_teacher(user)
    if credentials is None:
        raise HTTPException(status_code=401, detail="未登录")
    try:
        teacher = _clear_workspace_data(db, user, load_seed=False)
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"重置失败: {exc}") from exc
    return {
        "ok": True,
        "teacher": {"name": teacher.name, "phone": teacher.phone},
    }


@router.post("/data/import/roster")
async def import_roster(
    file: UploadFile = File(...),
    class_id: uuid_mod.UUID | None = Form(default=None),
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Import student rows from xlsx. Updates existing students by admission_no.
    Optional class_id assigns imported rows to that class; otherwise new students
    go to the system unassigned class and existing enrollments are left alone."""
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析该 Excel 文件，请提供 .xlsx 格式")
    try:
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    header_idx, columns = _locate_header(rows)
    ignored = sorted(
        {_cell_str(c) for row in rows[header_idx : header_idx + 1] for c in row if _cell_str(c)}
        - set().union(*_HEADER_ALIASES.values())
    )
    score_columns = [c for c in ignored if _SCORE_COLUMN_HINT.search(c)]
    if score_columns:
        shown = "、".join(f"「{c}」" for c in score_columns[:3])
        raise HTTPException(
            status_code=400,
            detail=f"这像是成绩表格而不是花名册（检测到列：{shown}）。请选择花名册文件（学号/姓名/性别）导入",
        )

    assign_class = class_id is not None
    target_cls = _resolve_target_class(db, class_id, user)
    today = dt.date.today()
    report: list[dict] = []
    seen_in_file: set[str] = set()
    created = updated = 0
    for pos, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        values = {
            field: (_cell_str(row[col]) if col < len(row) else "")
            for field, col in columns.items()
        }
        admission_no = values.get("admission_no", "")
        name = values.get("name", "")
        if not admission_no and not name:
            continue
        item = {"row": pos, "admission_no": admission_no, "name": name, "status": "ok"}
        try:
            if not admission_no:
                raise ValueError("缺少学号")
            if not name:
                raise ValueError("缺少姓名")
            if admission_no in seen_in_file:
                raise ValueError("文件内学号重复")
            seen_in_file.add(admission_no)

            gender_raw = values.get("gender") or None
            gender = parse_gender(gender_raw) if gender_raw else None
            birth_raw = (
                row[columns["birth_date"]]
                if "birth_date" in columns and columns["birth_date"] < len(row)
                else None
            )
            try:
                birth_date = _parse_birth_date(birth_raw)
                item["message"] = None
            except ValueError as exc:
                birth_date, item["message"] = None, str(exc)
            address = (values.get("address") or "").strip() or None
            guardian_name = (values.get("guardian_name") or "").strip()
            guardian_phone = (values.get("guardian_phone") or "").strip() or None
            guardian_relation = (values.get("guardian_relation") or "").strip() or None

            person_row = (
                db.query(Person)
                .filter(
                    Person.payload["role"].as_string() == "student",
                    Person.payload["admission_no"].as_string() == admission_no,
                )
                .first()
            )
            if person_row is None:
                payload = validate_person_payload(
                    "student",
                    {
                        "admission_no": admission_no,
                        "gender": gender,
                        "birth_date": birth_date.isoformat() if birth_date else None,
                        "address": address,
                    },
                )
                person_row = Person(
                    name=name,
                    password_hash=hash_password(uuid_mod.uuid4().hex),
                    payload=payload,
                )
                db.add(person_row)
                db.flush()
                tag_student_workspace(person_row, user)
                _move_student(db, person_row, target_cls, reason="admitted", today=today)
                create_event(
                    db,
                    event_type="enrolled",
                    title="入学",
                    start_time=utcnow(),
                    payload=(
                        {} if is_unassigned_class(target_cls)
                        else {"class_name": target_cls.name}
                    ),
                    attendee_ids=[person_row.id],
                )
                created += 1
                item["status"] = "created"
            else:
                # 更新 = 按学号匹配已有学生并覆盖资料；把改了什么写进报告，
                # 否则「更新 N 人」对老师来说等于黑箱
                changes: list[str] = []
                old_payload = dict(person_row.payload or {})
                if person_row.name != name:
                    changes.append(f"姓名 {person_row.name} → {name}")
                person_row.name = name
                if gender is not None and old_payload.get("gender") != gender:
                    changes.append("性别")
                if birth_date and old_payload.get("birth_date") != birth_date.isoformat():
                    changes.append("出生日期")
                if address and old_payload.get("address") != address:
                    changes.append("家庭住址")
                payload = old_payload
                if gender is not None:
                    payload["gender"] = gender
                if birth_date:
                    payload["birth_date"] = birth_date.isoformat()
                if address:
                    payload["address"] = address
                person_row.payload = validate_person_payload("student", payload)
                if assign_class:
                    _move_student(db, person_row, target_cls, reason="moved", today=today)
                updated += 1
                item["status"] = "updated"
                if changes:
                    item["changes"] = "；".join(changes)

            sync_birthday_event(db, person_row)

            # 监护人列：给了姓名或电话就把该监护人挂到学生上（按电话/姓名合并），
            # 已有的其他监护人不改动；三列都为空则完全跳过
            if guardian_name or guardian_phone:
                guardian = _find_or_create_guardian(db, guardian_name, guardian_phone)
                _guardian_link(db, person_row.id, guardian, guardian_relation)
                item["guardian"] = guardian.name
        except ValueError as exc:
            item["status"] = "error"
            item["message"] = str(exc)
        report.append(item)

    errors = [r for r in report if r["status"] == "error"]
    db.commit()
    return {
        "target_class": class_for_api(target_cls),
        "created": created,
        "updated": updated,
        "errors": errors,
        "ignored_columns": ignored,
        "rows": report,
    }


def _roster_workbook(db: Session, cls: Class, students: list[Person]) -> Workbook:
    """导出与导入同格式的花名册：学号/姓名/性别/出生日期/家庭住址 + 首位
    监护人（姓名/电话/关系）；一名学生有多位监护人时只导出第一位。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "花名册"
    ws.append([class_export_heading(cls, "花名册")])
    ws.append(_ROSTER_HEADERS)
    for s in students:
        payload = s.payload or {}
        guardians = _guardians_of(db, s.id)
        guardian, relationship = guardians[0] if guardians else (None, None)
        g_payload = guardian.payload or {} if guardian else {}
        ws.append([
            payload.get("admission_no") or "",
            s.name,
            gender_label(payload.get("gender")),
            payload.get("birth_date") or "",
            payload.get("address") or "",
            guardian.name if guardian else "",
            g_payload.get("phone") or "",
            relationship or "",
        ])
    widths = {"A": 16, "B": 12, "C": 8, "D": 14, "E": 24, "F": 14, "G": 16, "H": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    bold = Font(bold=True)
    ws.cell(row=1, column=1).font = bold
    for col in range(1, len(_ROSTER_HEADERS) + 1):
        ws.cell(row=2, column=col).font = bold
    return wb


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    quoted = quote(filename)
    return StreamingResponse(
        buf,
        media_type=_XLSX_MIME,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"},
    )


@router.get("/data/export/roster")
def export_roster(
    class_id: uuid_mod.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Export one class's current roster in the import-compatible format."""
    cls = _class_for_export(db, user, class_id)
    students = _class_roster_students(db, cls)
    return _xlsx_response(
        _roster_workbook(db, cls, students),
        f"{class_export_heading(cls, '花名册')}.xlsx",
    )


def _home_visits_workbook(
    db: Session,
    cls: Class,
    students: list[Person],
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "家访记录"
    ws.append([class_export_heading(cls, "家访记录")])
    ws.append(_HOME_VISIT_HEADERS)
    student_ids = [s.id for s in students]
    students_by_id = {s.id: s for s in students}
    if student_ids:
        rows = (
            db.query(Event, person_events.c.person_id)
            .join(person_events, person_events.c.event_id == Event.id)
            .filter(
                Event.type == "home_visited",
                person_events.c.person_id.in_(student_ids),
            )
            .order_by(Event.start_time.desc(), Event.created_at.desc())
            .all()
        )
        seen: set[tuple] = set()
        for ev, sid in rows:
            if sid not in students_by_id:
                continue
            key = (ev.id, sid)
            if key in seen:
                continue
            seen.add(key)
            student = students_by_id[sid]
            payload = ev.payload or {}
            done = payload.get("done")
            ws.append([
                (student.payload or {}).get("admission_no") or "",
                student.name,
                ev.start_time.strftime("%Y-%m-%d %H:%M"),
                payload.get("purpose") or "",
                payload.get("summary") or ev.title or "",
                payload.get("guardian") or "",
                "已完成" if done else "未完成",
            ])
    widths = {"A": 16, "B": 12, "C": 18, "D": 14, "E": 28, "F": 12, "G": 10}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    bold = Font(bold=True)
    ws.cell(row=1, column=1).font = bold
    for col in range(1, len(_HOME_VISIT_HEADERS) + 1):
        ws.cell(row=2, column=col).font = bold
    return wb


def _scores_workbook(
    db: Session,
    cls: Class,
    students: list[Person],
    user: Person,
) -> Workbook:
    wid = workspace_id(user)
    student_ids = {s.id for s in students}
    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()
    sheet_count = 0
    for exam in exam_events(db, wid):
        if not subjects_config(exam):
            continue
        if not exam_relevant_to_class(db, exam, student_ids, wid):
            continue
        ws = wb.create_sheet(unique_sheet_title(exam.title, used_titles))
        scores_map = scores_for_students(db, exam, list(student_ids), wid)
        append_score_sheet(ws, exam, students, scores_map)
        sheet_count += 1
    if sheet_count == 0:
        raise HTTPException(status_code=400, detail="该班暂无考试成绩可导出")
    return wb


@router.get("/data/export/home-visits")
def export_home_visits(
    class_id: uuid_mod.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Export home-visit records for students currently enrolled in one class."""
    cls = _class_for_export(db, user, class_id)
    students = _class_roster_students(db, cls)
    return _xlsx_response(
        _home_visits_workbook(db, cls, students),
        f"{class_export_heading(cls, '家访记录')}.xlsx",
    )


@router.get("/data/export/scores")
def export_scores(
    class_id: uuid_mod.UUID,
    db: Session = Depends(get_db),
    user: Person = Depends(get_current_person),
):
    """Export class scores as one workbook with one sheet per relevant exam."""
    cls = _class_for_export(db, user, class_id)
    students = _class_roster_students(db, cls)
    return _xlsx_response(
        _scores_workbook(db, cls, students, user),
        f"{class_export_heading(cls, '成绩')}.xlsx",
    )


@router.get("/data/export/roster-template")
def export_roster_template():
    """Blank roster template in the import-compatible format."""
    wb = Workbook()
    ws = wb.active
    ws.title = "花名册"
    ws.append(["班级花名册"])
    ws.append(_ROSTER_HEADERS)
    ws.append([
        "2025070701", "张三", "男", "2012-05-14", "解放路100号",
        "张丽", "13900000001", "母亲",
    ])
    ws.append([
        "2025070702", "李四", "女", "2012-11-02", "解放路101号",
        "李强", "13900000002", "父亲",
    ])
    return _xlsx_response(wb, "花名册模板.xlsx")
